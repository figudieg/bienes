import openpyxl
import pandas as pd
import datetime
from django.db import transaction
from django.contrib.auth import get_user_model
from .models import Bien, Sede, Area, TrazabilidadMovimientos, MantenimientoBien, OrdenCompra, Funcionario

User = get_user_model()

def clean_val(val):
    if val is None:
        return ""
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    return str(val).strip()

def parse_date(val):
    if isinstance(val, (datetime.date, datetime.datetime)):
        return val.date() if isinstance(val, datetime.datetime) else val
    if not val:
        return datetime.date.today()
    s = str(val).strip()
    for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d/%m/%y', '%d-%m-%Y', '%d%m%Y'):
        try:
            return datetime.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return datetime.date.today()

def clean_serial(serial_val):
    s = clean_val(serial_val)
    if not s or s.upper() in ["S/S", "S/N", "S.S.", "SIN SERIAL", "N/A", "NONE", "NAN", "-"]:
        return None
    return s

def is_invalid_bien_code(code_val):
    c = clean_val(code_val).upper()
    if not c or c in [
        "NONE", "NAN", "", "0", "N° DE BIEN", "N°", "RESPONSABLE", 
        "RESPONSABLE ADMINISTRATIVO CEDENTE", "RESPONSABLE ADMINISTRATIVO RECEPTOR",
        "FIRMA Y SELLO", "C.I. N°", "ORGANISMO", "DENOMINACION", 
        "CÓDIGO SIGECOFF", "CODIGO SIGECOFF", "TOTAL", "TOTAL BIENES", "BIENES TOTALES",
        "ELABORADO POR", "NOMBRE Y APELLIDO", "CARGO", "FIRMA", "NOTA:", "S/B"
    ] or "TOTAL" in c or "FIRMA" in c or "ELABORADO" in c or "REPARADO" in c or "CONFORMADO" in c:
        return True
    return False

def check_serial_uniqueness(serial, nro_bien, results, row_num):
    if not serial:
        return None
    existing_bien = Bien.objects.filter(serial_fabrica=serial).exclude(codigo_inventario=nro_bien).first()
    if existing_bien:
        results['errors'].append(
            f"Fila {row_num}: El serial '{serial}' ya lo tiene el bien '{existing_bien.codigo_inventario}'. "
            f"Se importa este registro sin serial."
        )
        return None
    return serial

def import_excel_file(file_path, file_name):
    results = {
        'status': 'success',
        'imported_count': 0,
        'errors': [],
        'type_detected': 'unknown'
    }

    lower_name = file_name.lower()
    
    try:
        default_sede, _ = Sede.objects.get_or_create(nombre="Sede Central DEM")
        default_area, _ = Area.objects.get_or_create(nombre="Dirección General", sede=default_sede)
        default_user = User.objects.filter(is_superuser=True).first()
        if not default_user:
            default_user = User.objects.create_superuser(
                username='system_import', email='import@bienespublicos.local', password='Password123!'
            )
        default_funcionario, _ = Funcionario.objects.get_or_create(
            cedula='V-00000000',
            defaults={
                'nombres': 'Importación',
                'apellidos': 'Masiva',
                'cargo': 'Sistema',
                'area': default_area
            }
        )

        if "incorporación" in lower_name or "incorporacion" in lower_name:
            results['type_detected'] = 'Comprobante de Incorporación'
            wb = openpyxl.load_workbook(file_path, data_only=True)
            sheet = wb.active
            
            nro_comp = clean_val(sheet.cell(row=2, column=8).value) or "INC-GEN"
            fecha_comp = parse_date(sheet.cell(row=4, column=8).value)
            
            oc, _ = OrdenCompra.objects.get_or_create(
                numero_orden=f"INC-{nro_comp}",
                defaults={
                    'proveedor': "Bienes Públicos / Incorporación Masiva",
                    'fecha_llegada': fecha_comp,
                    'conformidad_recepcion': True
                }
            )

            with transaction.atomic():
                for r in range(15, sheet.max_row + 1):
                    nro_bien = clean_val(sheet.cell(row=r, column=1).value)
                    if is_invalid_bien_code(nro_bien):
                        continue
                    
                    descripcion = clean_val(sheet.cell(row=r, column=2).value) or "Sin descripción"
                    marca = clean_val(sheet.cell(row=r, column=3).value)
                    modelo = clean_val(sheet.cell(row=r, column=4).value)
                    serial = clean_serial(sheet.cell(row=r, column=5).value)
                    serial = check_serial_uniqueness(serial, nro_bien, results, r)
                    
                    tipo_inc = clean_val(sheet.cell(row=r, column=6).value)
                    condicion = clean_val(sheet.cell(row=r, column=7).value)
                    
                    bien, created = Bien.objects.update_or_create(
                        codigo_inventario=nro_bien,
                        defaults={
                            'nombre': descripcion[:150],
                            'descripcion': f"{descripcion}. Marca: {marca}, Modelo: {modelo}, Condición: {condicion}, Tipo Inc: {tipo_inc}",
                            'serial_fabrica': serial,
                            'sede': default_sede,
                            'orden_compra': oc,
                            'estado': 'ACTIVO'
                        }
                    )
                    
                    TrazabilidadMovimientos.objects.create(
                        bien=bien,
                        tipo_movimiento='INCORPORACION',
                        sede_destino=default_sede,
                        area_destino=default_area,
                        funcionario_destino=default_funcionario,
                        motivo=f"Incorporación según comprobante {nro_comp}",
                        usuario_responsable=default_user
                    )
                    results['imported_count'] += 1

        elif "reasignación" in lower_name or "reasignacion" in lower_name:
            results['type_detected'] = 'Comprobante de Reasignación'
            wb = openpyxl.load_workbook(file_path, data_only=True)
            
            with transaction.atomic():
                for name in wb.sheetnames:
                    sheet = wb[name]
                    nro_comp = clean_val(sheet.cell(row=2, column=7).value) or "REAS-GEN"
                    
                    for r in range(15, sheet.max_row + 1):
                        nro_bien = clean_val(sheet.cell(row=r, column=1).value)
                        if is_invalid_bien_code(nro_bien):
                            continue
                        
                        try:
                            bien = Bien.objects.get(codigo_inventario=nro_bien)
                            
                            TrazabilidadMovimientos.objects.create(
                                bien=bien,
                                tipo_movimiento='REASIGNACION',
                                sede_origen=bien.sede,
                                sede_destino=default_sede,
                                area_destino=default_area,
                                funcionario_destino=default_funcionario,
                                motivo=f"Reasignación según comprobante {nro_comp} (Hoja: {name})",
                                usuario_responsable=default_user
                            )
                            results['imported_count'] += 1
                        except Bien.DoesNotExist:
                            results['errors'].append(f"Fila {r} ({name}): El bien con código {nro_bien} no existe en el sistema. Debe incorporarse primero.")

        elif "mantenimiento" in lower_name:
            results['type_detected'] = 'Ficha de Mantenimiento'
            wb = openpyxl.load_workbook(file_path, data_only=True)
            
            with transaction.atomic():
                for name in wb.sheetnames:
                    sheet = wb[name]
                    nro_ficha = clean_val(sheet.cell(row=2, column=10).value) or name
                    fecha_f = parse_date(sheet.cell(row=3, column=10).value)
                    
                    for r in range(13, sheet.max_row + 1):
                        especif = clean_val(sheet.cell(row=r, column=1).value)
                        nro_bien = clean_val(sheet.cell(row=r, column=4).value)
                        if is_invalid_bien_code(nro_bien) or especif == "Reparado por:":
                            continue
                        
                        tipo_m = clean_val(sheet.cell(row=r, column=5).value).upper()
                        if "PREVENTIVO" in tipo_m:
                            tipo_m = 'PREVENTIVO'
                        else:
                            tipo_m = 'CORRECTIVO'
                            
                        actividad = clean_val(sheet.cell(row=r, column=6).value) or "Mantenimiento General"
                        materiales = clean_val(sheet.cell(row=r, column=7).value)
                        factura = clean_val(sheet.cell(row=r, column=8).value) or "Autogestión"
                        
                        costo_val = sheet.cell(row=r, column=9).value
                        costo = abs(float(costo_val)) if (costo_val is not None and str(costo_val).strip() != "") else 0.0
                        
                        fecha_m = parse_date(sheet.cell(row=r, column=10).value)
                        nota = clean_val(sheet.cell(row=r, column=11).value)
                        
                        bien, _ = Bien.objects.get_or_create(
                            codigo_inventario=nro_bien,
                            defaults={
                                'nombre': especif[:150],
                                'descripcion': especif,
                                'sede': default_sede,
                                'estado': 'ACTIVO'
                            }
                        )
                        
                        MantenimientoBien.objects.create(
                            bien=bien,
                            numero_ficha=nro_ficha,
                            fecha_ficha=fecha_f,
                            tipo_mantenimiento=tipo_m,
                            actividad_realizada=actividad,
                            materiales_empleados=materiales,
                            numero_factura=factura,
                            costo=costo,
                            fecha_mantenimiento=fecha_m,
                            nota=nota,
                            reparado_por="Rafael Jaimes",
                            conformado_por="Yohaly Alarcón de Marcano",
                            responsable_administrativo="Responsable de Bienes"
                        )
                        results['imported_count'] += 1

        elif "inventario" in lower_name:
            results['type_detected'] = 'Inventario Bienes Muebles'
            
            if file_name.endswith('.xls'):
                df = pd.read_excel(file_path, header=None)
                with transaction.atomic():
                    for idx, row in df.iloc[4:].iterrows():
                        nro_bien = clean_val(row[0])
                        if is_invalid_bien_code(nro_bien):
                            continue
                        
                        desc = clean_val(row[1]) or "Bienes Muebles"
                        marca = clean_val(row[2])
                        modelo = clean_val(row[3])
                        serial = clean_serial(row[4])
                        serial = check_serial_uniqueness(serial, nro_bien, results, idx+1)
                        tipo_clase = clean_val(row[5])
                        
                        bien, _ = Bien.objects.update_or_create(
                            codigo_inventario=nro_bien,
                            defaults={
                                'nombre': desc[:150],
                                'descripcion': f"{desc}. Marca: {marca}, Modelo: {modelo}, Tipo: {tipo_clase}",
                                'serial_fabrica': serial,
                                'sede': default_sede,
                                'estado': 'ACTIVO'
                            }
                        )
                        results['imported_count'] += 1
            else:
                wb = openpyxl.load_workbook(file_path, data_only=True)
                sheet = wb.active
                with transaction.atomic():
                    for r in range(15, sheet.max_row + 1):
                        nro_bien = clean_val(sheet.cell(row=r, column=1).value)
                        if is_invalid_bien_code(nro_bien):
                            continue
                        
                        desc = clean_val(sheet.cell(row=r, column=2).value) or "Bienes Muebles"
                        marca = clean_val(sheet.cell(row=r, column=4).value)
                        modelo = clean_val(sheet.cell(row=r, column=5).value)
                        serial = clean_serial(sheet.cell(row=r, column=6).value)
                        serial = check_serial_uniqueness(serial, nro_bien, results, r)
                        tipo_clase = clean_val(sheet.cell(row=r, column=11).value)
                        
                        bien, _ = Bien.objects.update_or_create(
                            codigo_inventario=nro_bien,
                            defaults={
                                'nombre': desc[:150],
                                'descripcion': f"{desc}. Marca: {marca}, Modelo: {modelo}, Tipo: {tipo_clase}",
                                'serial_fabrica': serial,
                                'sede': default_sede,
                                'estado': 'ACTIVO'
                            }
                        )
                        results['imported_count'] += 1
        else:
            results['status'] = 'error'
            results['errors'].append("No se pudo clasificar el tipo de archivo Excel. Verifique el nombre del archivo.")

    except Exception as e:
        results['status'] = 'error'
        results['errors'].append(f"Excepción general durante la importación: {str(e)}")
        
    return results
